# -*- coding: utf-8 -*-
"""扱える形式の回帰テスト（.docx / .xlsx / .doc / .xls / PDF）。

PDFは3つの型を見分けられることを確かめる。画像PDFを「埋められない」と正しく言えることが大事で、
当て推量で書くと申請書として事故になる。

  .venv/bin/python -m pytest tests/ -q
"""
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import legacy                                    # noqa: E402
from app.forms import find_blanks_docx, find_blanks_xlsx  # noqa: E402
from app.pdfform import apply_pdf, find_blanks_pdf, kind_of  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAMPLE = os.path.join(ROOT, "samples", "J-SHIS利用申請（参考例）.docx")


def _make_xls(path):
    import xlwt
    w = xlwt.Workbook()
    s = w.add_sheet("申請")
    s.write(0, 0, "申請日"); s.write(0, 1, "令和○年○月○日")
    s.write(1, 0, "事業者名"); s.write(1, 1, "〇〇株式会社")
    w.save(path)


def _make_acroform(path):
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(path)
    c.drawString(60, 760, "test form")
    c.acroForm.textfield(name="company", x=150, y=700, width=300, height=22,
                         borderWidth=1, fontName="Helvetica", fontSize=11)
    c.save()


def _make_flat_pdf(path):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas
    pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
    c = canvas.Canvas(path)
    c.setFont("HeiseiKakuGo-W5", 12)
    c.drawString(60, 700, "利用申請書")
    c.drawString(60, 670, "事業者名 〇〇株式会社")
    c.drawString(60, 640, "申請日 令和○年○月○日")
    c.save()


def test_docx_blanks():
    blanks, _ = find_blanks_docx(SAMPLE)
    assert len(blanks) == 5


def test_xlsx_and_xls(tmp_path):
    xls = str(tmp_path / "t.xls")
    _make_xls(xls)
    assert legacy.available(), "LibreOffice が要る（apt install libreoffice-writer libreoffice-calc）"
    xlsx = legacy.convert(xls, ".xlsx", str(tmp_path / "conv"))
    blanks, _ = find_blanks_xlsx(xlsx)
    raws = [b.raw for b in blanks]
    assert any("年" in r for r in raws), raws
    assert any(r in ("〇〇", "○○") for r in raws), raws
    # 元の形式へ戻せること（指定様式のまま出すのが要件）
    back = legacy.convert(xlsx, ".xls", str(tmp_path / "back"))
    assert os.path.getsize(back) > 0


def test_doc_roundtrip(tmp_path):
    assert legacy.available()
    doc = legacy.convert(SAMPLE, ".doc", str(tmp_path / "d"))
    back = legacy.convert(doc, ".docx", str(tmp_path / "b"))
    a = [p for p in _texts(SAMPLE)]
    b = [p for p in _texts(back)]
    assert a == b, "往復で文字が変わった"


def _texts(path):
    import docx
    return [p.text for p in docx.Document(path).paragraphs if p.text.strip()]


def test_pdf_kinds(tmp_path):
    acro = str(tmp_path / "a.pdf"); _make_acroform(acro)
    flat = str(tmp_path / "f.pdf"); _make_flat_pdf(flat)
    assert kind_of(acro) == "acroform"
    assert kind_of(flat) == "flat"
    bs, k = find_blanks_pdf(acro)
    assert k == "acroform" and [b.locator for b in bs] == ["company"]
    bs, k = find_blanks_pdf(flat)
    assert k == "flat" and any(b.raw in ("〇〇", "○○") for b in bs)


def test_pdf_acroform_fill(tmp_path):
    import pypdf
    acro = str(tmp_path / "a.pdf"); _make_acroform(acro)
    bs, k = find_blanks_pdf(acro)
    bs[0].value = "株式会社エクスブリッジ"
    out = str(tmp_path / "filled.pdf")
    apply_pdf(acro, out, bs, k)
    r = pypdf.PdfReader(out)
    assert r.get_fields()["company"].get("/V") == "株式会社エクスブリッジ"
    # 欄のフォントが日本語を持たないPDFでも化けないよう、ビューアに描かせる
    # pypdf は BooleanObject を返すので `is True` では比較できない
    assert bool(r.trailer["/Root"]["/AcroForm"].get("/NeedAppearances")) is True


def test_image_pdf_is_refused(tmp_path):
    """画像PDFは『埋められない』と返すこと。当て推量で書かない。"""
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    import io
    from PIL import Image
    img = Image.new("RGB", (400, 300), "white")
    buf = io.BytesIO(); img.save(buf, format="PNG"); buf.seek(0)
    p = str(tmp_path / "img.pdf")
    c = canvas.Canvas(p)
    c.drawImage(ImageReader(buf), 50, 400, width=400, height=300)
    c.save()
    bs, k = find_blanks_pdf(p)
    assert k == "image" and bs == []
